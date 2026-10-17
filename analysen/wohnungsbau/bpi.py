"""Baupreisindex des BFS (Basis Oktober 2020 = 100) als Jahresmittel aus April und Oktober.

Quelle: `quellen/baupreisindex-cc-t-05.05.01.xlsx` (BFS, cc-t-05.05.01, Multibasen), Blatt «2020».
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

QUELLE = Path(__file__).resolve().parent / "quellen" / "baupreisindex-cc-t-05.05.01.xlsx"
OBJEKTE = {"<OBJ_05>": "neubau_mfh", "<OBJ_11>": "renovation_mfh"}


def jahresindex() -> pd.DataFrame:
    """Je Grossregion (und Schweiz) und Jahr: Index Neubau und Renovation Mehrfamilienhaus."""
    wb = openpyxl.load_workbook(QUELLE, read_only=True)
    beschriftung = {r[0]: r[4] for r in wb["BAP"].iter_rows(values_only=True) if r[0]}
    zeilen = list(wb["2020"].iter_rows(values_only=True))
    jahre = zeilen[5]
    region, werte = None, []
    for z in zeilen[7:]:
        if z[0] and str(z[0]).startswith("<REG"):
            region = beschriftung[z[0]].split(" (")[0]
        elif z[0] in OBJEKTE and region:
            for jahr, wert in zip(jahre[3:], z[3:]):
                if jahr and isinstance(wert, (int, float)):
                    werte.append((region, int(jahr), OBJEKTE[z[0]], float(wert)))
    d = pd.DataFrame(werte, columns=["region", "jahr", "objekt", "index"])
    # Nur Jahre mit beiden Erhebungen (April und Oktober) zählen als Jahresmittel.
    g = d.groupby(["region", "jahr", "objekt"])["index"].agg(["mean", "size"]).reset_index()
    g = g[g["size"] == 2].drop(columns="size")
    return g.pivot_table(index=["region", "jahr"], columns="objekt", values="mean").reset_index()


if __name__ == "__main__":
    print(jahresindex().query("region == 'Schweiz' and jahr >= 2010").round(1).to_string())
